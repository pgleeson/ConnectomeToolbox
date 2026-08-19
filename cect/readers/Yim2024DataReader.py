# -*- coding: utf-8 -*-

############################################################

#    A script to read the values of Yim et al. 2024

############################################################


from cect.ConnectomeReader import ConnectionInfo
from cect.Cells import convert_to_preferred_muscle_name
from cect.Cells import is_any_neuron
from cect.Cells import remove_leading_index_zero
from cect.Cells import is_potential_muscle
from cect.Cells import is_known_muscle
from cect.ConnectomeDataset import get_dataset_source_on_github
from cect.ConnectomeDataset import LOAD_READERS_FROM_CACHE_BY_DEFAULT

from cect.ConnectomeDataset import ConnectomeDataset

from cect.Neurotransmitters import CONTACTOME_SYN_TYPE
from cect.Neurotransmitters import CONTACTOME_SYN_CLASS

from cect.Neurotransmitters import CHEMICAL_SYN_TYPE
from cect.Neurotransmitters import GENERIC_CHEM_SYN_CLASS

from openpyxl import load_workbook

import os
import numpy as np

from cect import print_

spreadsheet_location = os.path.dirname(os.path.abspath(__file__)) + "/../data/"

SYNAPTIC_CONNS_FILENAME = "%s41467_2024_45943_MOESM6_ESM.xlsx" % spreadsheet_location
CONTACTOME_FILENAME = "%s41467_2024_45943_MOESM9_ESM.xlsx" % spreadsheet_location

DAUER_NON_NORM = "Dauer"
DAUER_NORM = "Dauer_normalized"


NAME = "Yim2024"

pre_range = range(3, 225)
post_range = range(3, 225)


def get_syntype_synclass(conn_filename):
    if conn_filename == SYNAPTIC_CONNS_FILENAME:
        return CHEMICAL_SYN_TYPE, GENERIC_CHEM_SYN_CLASS
    else:
        return CONTACTOME_SYN_TYPE, CONTACTOME_SYN_CLASS


READER_DESCRIPTION = (
    """Data extracted from %s Yim et al. 2024 on Dauer connectome (Synaptic connections; Normalized)"""
    % get_dataset_source_on_github(SYNAPTIC_CONNS_FILENAME.split("/")[-1])
)

DATASET_DESCRIPTION_0 = """Reconstruction of the directed chemical synaptic connectome of the dauer, a distinct developmental stage of _C. elegans_. Presynaptic active zones were detected by a convolutional neural network, reconstructed in 3D and proofread; postsynaptic partners and each partner's share of an active zone were assigned by simulating neurotransmitter diffusion. """

DATASET_DESCRIPTION = (
    DATASET_DESCRIPTION_0
    + """This connectome dataset contains normalized weights to ease comparison to other datasets."""
)

WEIGHTS_0 = "Weights are the summed volume (nm<sup>3</sup>) of active zone material attributed to a pre/post pair."
WEIGHTS = (
    WEIGHTS_0
    + ". In this dataset, these are normalized by the standard deviation of connection weights without the top 5th percentile to remove the bias due to the big outliers"
)


class Yim2024DataReader(ConnectomeDataset):
    """
    Reader of data from Yim et al. 2024 - Dauer connectome

    Returns:
        (Yim2024DataReader): The initialized connectome reader
    """

    verbose = False

    def __init__(self, normalized, conn_filename):

        ConnectomeDataset.__init__(self)

        conn_type = DAUER_NORM if normalized else DAUER_NON_NORM
        self.conn_filename = conn_filename

        print_(f"Opening sheet {conn_type} in the Excel file: {conn_filename}")

        wb = load_workbook(conn_filename)

        self.pre_cells = {}
        self.post_cells = {}
        self.conn_nums = {}
        self.normalized = normalized

        sheet = wb.get_sheet_by_name(conn_type)

        self.pre_cells[conn_type] = []
        self.post_cells[conn_type] = []

        for i in pre_range:
            self.pre_cells[conn_type].append(sheet["A%i" % i].value)

        if self.verbose:
            print_(
                " - Pre cells for %s (%i):\n%s"
                % (
                    conn_type,
                    len(self.pre_cells[conn_type]),
                    self.pre_cells[conn_type],
                )
            )

        for i in post_range:
            self.post_cells[conn_type].append(sheet.cell(row=1, column=i).value)

        if self.verbose:
            print_(
                " - Post cells for %s (%i):\n%s"
                % (
                    conn_type,
                    len(self.post_cells[conn_type]),
                    self.post_cells[conn_type],
                )
            )

        self.conn_nums[conn_type] = np.zeros(
            [len(self.pre_cells[conn_type]), len(self.post_cells[conn_type])],
            dtype=float,
        )

        for i in range(len(self.pre_cells[conn_type])):
            for j in range(len(self.post_cells[conn_type])):
                row = 3 + i
                col = 3 + j
                val = sheet.cell(row=row, column=col).value
                if val != 0:
                    print_("Cell (%i,%i) [row %i, col %i] = %s" % (i, j, row, col, val))
                if val is not None:
                    self.conn_nums[conn_type][i, j] = val

        if self.verbose:
            print_(
                " - Conns for %s (%s):\n%s"
                % (
                    conn_type,
                    self.conn_nums[conn_type].shape,
                    self.conn_nums[conn_type],
                )
            )

        neurons, muscles, other_cells, conns = self.read_all_data()

        for conn in conns:
            self.add_connection_info(conn)

    def read_data(self):
        return self._read_data()

    def read_muscle_data(self):
        return self._read_muscle_data()

    def read_all_data(self):
        """
        Returns:
            (tuple[list, list, list, list]): List of neurons, muscles, other cells and connections which have been read in
        """

        neurons = set([])
        muscles = set([])
        other_cells = set([])
        conns = []

        conn_type = DAUER_NORM if self.normalized else DAUER_NON_NORM

        for pre_index in range(len(self.pre_cells[conn_type])):
            for post_index in range(len(self.post_cells[conn_type])):
                num = self.conn_nums[conn_type][pre_index, post_index]

                pre = remove_leading_index_zero(self.pre_cells[conn_type][pre_index])
                post = remove_leading_index_zero(self.post_cells[conn_type][post_index])
                if self.verbose and num > 0:
                    print_("Conn %s -> %s #%i" % (pre, post, num))

                if is_potential_muscle(pre):
                    pre = convert_to_preferred_muscle_name(pre)

                if is_potential_muscle(post):
                    post = convert_to_preferred_muscle_name(post)

                if num > 0:
                    syntype, synclass = get_syntype_synclass(self.conn_filename)

                    ci = ConnectionInfo(pre, post, num, syntype, synclass)
                    if self.verbose:
                        print_("Conn: %s" % (ci))
                    conns.append(ci)

                    for p in [pre, post]:
                        if is_any_neuron(p):
                            neurons.add(pre)
                        elif is_known_muscle(p):
                            muscles.add(pre)
                        else:
                            other_cells.add(p)

        return list(neurons), list(muscles), list(other_cells), conns


def get_instance(from_cache=LOAD_READERS_FROM_CACHE_BY_DEFAULT):
    """Uses ``Yim2024DataReader`` to load data on dauer connectome

    Returns:
        (Yim2024DataReader): The initialized connectome reader
    """
    if from_cache:
        from cect.ConnectomeDataset import (
            load_connectome_dataset_file,
            get_cache_filename,
        )

        return load_connectome_dataset_file(
            get_cache_filename(__file__.split("/")[-1].split(".")[0])
        )
    else:
        return Yim2024DataReader(normalized=True, conn_filename=SYNAPTIC_CONNS_FILENAME)


def main():
    tdr_instance = get_instance(from_cache=False)

    # analyse_connections(cells, neuron_conns, neurons2muscles, muscles, muscle_conns)

    print_(" -- Finished analysing connections using: %s" % os.path.basename(__file__))

    print(tdr_instance.summary())

    from cect.ConnectomeView import RAW_VIEW as view
    # from cect.ConnectomeView import PHARYNX_VIEW as view
    # from cect.ConnectomeView import NEURONS_VIEW as view

    print("=======================")
    cds2 = tdr_instance.get_connectome_view(view)
    print(cds2.summary(list_pre_cells=False))

    print("Plotting view: %s" % view)
    fig, _ = cds2.to_plotly_matrix_fig(
        "Chemical",
        view,
    )
    import plotly.io as pio

    pio.renderers.default = "browser"
    import sys

    if "-nogui" not in sys.argv:
        fig.show()


if __name__ == "__main__":
    main()
